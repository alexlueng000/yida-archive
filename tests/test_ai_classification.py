"""正文提取、模型协议、缓存和归档集成回归，不使用真实密钥。"""
import copy
import csv
import hashlib
import io
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from yida_archive import ai_classify as ai, classify, document_text, pipeline
from yida_archive.config import Config, load
from yida_archive.models import AttachmentObs, Instance
from yida_archive import dedupe


TEXT = 'Packing List\n正式装箱单，货物数量10件，净重100千克，毛重120千克。'
SECTIONS = [{'location': '段落1', 'text': TEXT}]


def decision(kind='箱单', status='classified', score=0.99):
    return {'status': status, 'type': kind, 'score': score, 'reason': '正文为正式装箱单并列出货物数量重量',
            'evidence': [{'location': '段落1', 'quote': '货物数量10件'}]}


def docx_bytes(text=TEXT):
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def files_for(name='资料.docx', data=b'body', form='合同审批'):
    o = AttachmentObs(form, 'LOCAL', 'I1', '', 1000, 'f', '附件', name, len(data),
                      'uuid', '', hashlib.sha256(data).hexdigest())
    return dedupe.group([Instance(form, 'LOCAL', 'I1', '', 1000, attachments=[o])])[0]


class AIClassificationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.cfg = Config({'ai': {'enabled': True, 'model': 'fake',
            'base_url': 'https://example.invalid/v1', 'threshold': 0.9},
            'storage': {'work_dir': self.tmp.name, 'out_dir': self.tmp.name},
            'scope': {'project_no': 'P1', 'project_secret': False}})
        self.env = patch.dict(os.environ, {'ARCHIVE_AI_API_KEY': 'secret-for-test'})
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def apply(self, name='资料.docx', result=None, data=b'body'):
        files = files_for(name, data)
        classify.classify(files, self.cfg)
        with patch.object(document_text, 'extract', return_value=SECTIONS), \
             patch.object(ai.Client, 'classify', return_value=(result or decision(), {'total_tokens': 30})) as call:
            ai.apply(files, self.cfg, lambda f: data)
        return files[0], call.call_count

    def test_content_classifies_generic_filename_and_reuses_cache(self):
        f, count = self.apply()
        self.assertEqual((f.type_name, f.folder, count), ('箱单', '清关送货', 1))
        f, count = self.apply()
        self.assertEqual(count, 0)
        self.assertTrue(f.ai_result['cached'])

    def test_rules_change_invalidates_cache(self):
        self.apply()
        self.cfg.raw['classification'] = {'rules': [{'folder': '清关送货', 'type': '箱单'}]}
        self.assertEqual(self.apply()[1], 1)

    def test_source_change_invalidates_cache(self):
        self.apply()
        f = files_for(form='货款支付')
        with patch.object(document_text, 'extract', return_value=SECTIONS), \
             patch.object(ai.Client, 'classify', return_value=(decision(), {})) as call:
            ai.apply(f, self.cfg, lambda f: b'body')
        self.assertEqual(call.call_count, 1)

    def test_low_score_not_archived(self):
        f, _ = self.apply(result=decision(score=0.5))
        self.assertFalse(f.folder)
        self.assertEqual(f.ai_result['status'], 'unknown')

    def test_hallucinated_type_or_evidence_rejected(self):
        for result in [decision('自行新增类型'), dict(decision(), evidence=[{'location': '第999页', 'quote': '货物数量10件'}]),
                       dict(decision(), evidence=[{'location': '段落1', 'quote': '不存在的文字'}]),
                       dict(decision(), score=True), dict(decision(), score=float('nan'))]:
            with self.subTest(result=result):
                f, _ = self.apply(result=result)
                self.assertFalse(f.folder)
                self.assertEqual(f.ai_result['status'], 'failed')

    def test_confirmed_filename_uses_rule_without_ai(self):
        f, count = self.apply(name='采购合同.docx')
        self.assertEqual(f.type_name, '采购合同')
        self.assertEqual(f.ai_result['status'], 'rule_classified')
        self.assertEqual(count, 0)

    def test_draft_exclusion_retained(self):
        f, _ = self.apply(name='箱单草稿.docx')
        self.assertFalse(f.folder)
        self.assertEqual(f.ai_result['status'], 'excluded')

    def test_unknown_mixed_excluded(self):
        for status in ['unknown', 'mixed', 'excluded']:
            f, _ = self.apply(name=status + '.docx', result=decision('', status))
            self.assertFalse(f.folder)

    def test_confirmed_filename_skips_ai_even_when_request_fails(self):
        f = files_for('采购合同.docx')
        classify.classify(f, self.cfg)
        with patch.object(document_text, 'extract', return_value=SECTIONS), \
             patch.object(ai.Client, 'classify', side_effect=ai.AIError('模型请求连接失败或超时')) as call:
            ai.apply(f, self.cfg, lambda f: b'body')
        self.assertEqual(f[0].type_name, '采购合同')
        self.assertEqual(f[0].ai_result['status'], 'rule_classified')
        call.assert_not_called()

    def test_failed_request_falls_back_to_suggested_rule(self):
        files = files_for('箱单2024.docx')
        classify.classify(files, self.cfg)
        self.assertEqual(files[0].confidence, classify.SUGGESTED)
        with patch.object(document_text, 'extract', return_value=SECTIONS), \
             patch.object(ai.Client, 'classify', side_effect=ai.AIError('模型请求连接失败或超时')) as call:
            ai.apply(files, self.cfg, lambda f: b'body')
        f = files[0]
        self.assertEqual((f.folder, f.type_name, f.confidence), ('清关送货', '箱单', classify.SUGGESTED))
        self.assertEqual(f.ai_result['status'], 'rule_fallback')
        self.assertIn('模型请求连接失败或超时', f.review_note)
        self.assertIn('未经正文核验', f.review_note)
        self.assertEqual(call.call_count, 1)

    def test_excluded_unknown_mixed_never_fall_back(self):
        for status in ['excluded', 'unknown', 'mixed']:
            with self.subTest(status=status):
                f, _ = self.apply(name='箱单2024.docx', result=decision('', status), data=status.encode())
                self.assertEqual((f.folder, f.type_name, f.confidence), ('', '', classify.MANUAL))
                self.assertEqual(f.ai_result['status'], status)

    def test_low_score_never_falls_back(self):
        f, _ = self.apply(name='箱单2024.docx', result=decision(score=0.5))
        self.assertFalse(f.folder)
        self.assertEqual(f.ai_result['status'], 'unknown')

    def test_fallback_only_for_suggested_not_manual(self):
        files = files_for('资料.docx')
        classify.classify(files, self.cfg)
        self.assertEqual(files[0].confidence, classify.MANUAL)
        with patch.object(document_text, 'extract', return_value=SECTIONS), \
             patch.object(ai.Client, 'classify', side_effect=ai.AIError('模型请求连接失败或超时')):
            ai.apply(files, self.cfg, lambda f: b'body')
        self.assertEqual((files[0].folder, files[0].type_name, files[0].confidence), ('', '', classify.MANUAL))
        self.assertEqual(files[0].ai_result['status'], 'failed')

    def test_failure_policy_pending_keeps_old_behavior(self):
        self.cfg.raw['ai']['failure_policy'] = 'pending'
        files = files_for('箱单2024.docx')
        classify.classify(files, self.cfg)
        self.assertEqual(files[0].confidence, classify.SUGGESTED)
        with patch.object(document_text, 'extract', return_value=SECTIONS), \
             patch.object(ai.Client, 'classify', side_effect=ai.AIError('模型不可用')):
            ai.apply(files, self.cfg, lambda f: b'body')
        self.assertEqual((files[0].folder, files[0].type_name, files[0].confidence), ('', '', classify.MANUAL))
        self.assertEqual(files[0].ai_result['status'], 'failed')

    def test_unsupported_format_falls_back(self):
        files = files_for('代理协议 23-50-10.doc')
        classify.classify(files, self.cfg)
        reader = Mock()
        with patch.object(ai, 'Client') as client:
            ai.apply(files, self.cfg, reader)
        f = files[0]
        self.assertEqual((f.folder, f.type_name, f.confidence), ('合同', '协议', classify.SUGGESTED))
        self.assertEqual(f.ai_result['status'], 'rule_fallback')
        self.assertIn('此格式正文', f.review_note)
        self.assertIn('未经正文核验', f.review_note)
        reader.assert_not_called()
        client.assert_not_called()

    def test_extraction_and_unexpected_failures_fall_back(self):
        for error, note in [(document_text.TextError('PDF没有文字层'), 'PDF没有文字层'),
                            (OSError('private-internal-detail'), 'AI识别失败')]:
            with self.subTest(error=error):
                files = files_for('彩页1-Sigma 560.pdf')
                classify.classify(files, self.cfg)
                with patch.object(document_text, 'extract', side_effect=error), \
                     patch.object(ai, 'Client') as client:
                    ai.apply(files, self.cfg, lambda f: b'body')
                f = files[0]
                self.assertEqual((f.folder, f.type_name, f.confidence), ('办证', '彩页说明书', classify.SUGGESTED))
                self.assertEqual(f.ai_result['status'], 'rule_fallback')
                self.assertIn(note, f.review_note)
                self.assertNotIn('private-internal-detail', f.review_note)
                client.assert_not_called()

    def test_removed_rule_cannot_be_restored(self):
        files = files_for('箱单2024.docx')
        classify.classify(files, self.cfg)
        self.assertEqual(files[0].confidence, classify.SUGGESTED)
        self.cfg.raw['classification'] = {'rules': [{'folder': '合同', 'type': '协议'}]}
        ai.apply(files, self.cfg, lambda f: None)
        self.assertFalse(files[0].folder)
        self.assertEqual(files[0].ai_result['status'], 'failed')

    def test_pipeline_fallback_written_and_reported_for_review(self):
        from yida_archive.fetch_local import load as local_load, LocalClient
        root = Path(self.tmp.name) / 'source'
        target = root / '批次' / '附件'
        target.mkdir(parents=True)
        source = target / '代理协议 23-50-10.doc'
        source.write_bytes(b'legacy-doc')
        with patch.object(ai, 'Client') as client:
            result = pipeline.run(self.cfg, local_load(root), client=LocalClient([root]),
                                  dry_run=False, verbose=False)
        client.assert_not_called()
        self.assertEqual(len(result.done), 1)
        self.assertEqual(Path(result.done[0][1]).read_bytes(), source.read_bytes())
        summary = classify.describe(result.files, self.cfg)
        self.assertEqual((summary['classified'], summary['manual'], summary['rule_fallback'], summary['ai_failed']),
                         (1, 0, 1, 0))
        with open(result.paths['csv'], encoding='utf-8-sig') as fh:
            row = next(csv.DictReader(fh))
        self.assertEqual(row['AI状态'], 'rule_fallback')
        self.assertEqual(row['确认度'], classify.SUGGESTED)
        self.assertIn('未经正文核验', row['备注'])
        markdown = Path(result.paths['markdown']).read_text(encoding='utf-8')
        self.assertIn('| 已归类 | 1 |\n| 正文未读、按文件名兜底 | 1 |', markdown)
        self.assertIn('rule_fallback', markdown)
        self.assertIn('未经正文核验', markdown)

    def test_images_never_read_or_sent(self):
        f = files_for('付款凭证.png')
        classify.classify(f, self.cfg)
        reader = Mock(side_effect=AssertionError('不应读取图片'))
        ai.apply(f, self.cfg, reader)
        self.assertEqual(f[0].type_name, '付款凭证')
        self.assertEqual(f[0].ai_result['status'], 'image_skipped')
        reader.assert_not_called()

    def test_manual_override_never_calls_ai(self):
        f = files_for()
        self.cfg.raw['file_reviews'] = {f[0].ident: {'type': '采购合同'}}
        classify.classify(f, self.cfg)
        reader = Mock()
        ai.apply(f, self.cfg, reader)
        self.assertEqual(f[0].type_name, '采购合同')
        reader.assert_not_called()

    def test_missing_content_and_hash_mismatch(self):
        for body in [None, b'different']:
            f = files_for('资料.docx')
            classify.classify(f, self.cfg)
            ai.apply(f, self.cfg, lambda f: body)
            self.assertFalse(f[0].folder)
            self.assertEqual(f[0].ai_result['status'], 'failed')

    def test_disabled_preserves_rules(self):
        self.cfg.raw['ai']['enabled'] = False
        f = files_for('采购合同.docx')
        classify.classify(f, self.cfg)
        ai.apply(f, self.cfg, Mock())
        self.assertEqual(f[0].type_name, '采购合同')

    def test_vertical_gbk_csv_and_relative_path(self):
        path = Path(self.tmp.name) / '凭证.csv'
        path.write_text('id,7509302\napiKey,test-key\nopenAiCompatible,https://example.invalid/v1\nworkspaceName,主账号空间\n', encoding='gb18030')
        with patch.dict(os.environ, {'ARCHIVE_AI_API_KEY': ''}):
            self.assertEqual(ai.credentials({'credentials_file': str(path), 'model': 'fake'})[0], 'test-key')
        config = Path(self.tmp.name) / 'config.yaml'
        config.write_text('ai:\n  credentials_file: 凭证.csv\n', encoding='utf-8')
        self.assertEqual(load(str(config)).raw['ai']['credentials_file'], str(path))

    def test_client_retries_and_never_exposes_response_or_key(self):
        client = ai.Client(dict(self.cfg.raw['ai'], retries=1))
        limited = Mock(status_code=429)
        success = Mock(status_code=200)
        success.json.return_value = {'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps(decision())}}], 'usage': {'total_tokens': 10}}
        with patch.object(ai.requests, 'post', side_effect=[limited, success]) as post, patch.object(ai.time, 'sleep'):
            result, _ = client.classify({'sections': SECTIONS})
            self.assertEqual(post.call_count, 2)
            self.assertFalse(post.call_args.kwargs['allow_redirects'])
            self.assertEqual(result['type'], '箱单')
        bad = Mock(status_code=401, text='secret-for-test')
        with patch.object(ai.requests, 'post', return_value=bad), self.assertRaises(ai.AIError) as caught:
            client.classify({})
        self.assertNotIn('secret-for-test', str(caught.exception))

    def test_horizontal_csv(self):
        path = Path(self.tmp.name) / 'key.csv'
        path.write_text('apiKey,base_url,model\ntest-key,https://example.invalid/v1,fake\n', encoding='utf-8')
        with patch.dict(os.environ, {'ARCHIVE_AI_API_KEY': ''}):
            self.assertEqual(ai.credentials({'credentials_file': str(path)}),
                             ('test-key', 'https://example.invalid/v1', 'fake'))

    def test_call_budget(self):
        client = ai.Client(dict(self.cfg.raw['ai'], max_calls=0))
        with patch.object(ai.requests, 'post') as post, self.assertRaises(ai.AIError):
            client.classify({})
        post.assert_not_called()

    def test_real_docx_xlsx_and_pdf_extraction(self):
        self.assertIn('货物数量10件', document_text.extract(docx_bytes(), '.docx')[0]['text'])
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = '箱单'
        book.active.append(['Packing List', '数量', '毛重'])
        book.active.append(['货物A', 10, 100])
        buf = io.BytesIO()
        book.save(buf)
        self.assertIn('工作表箱单', document_text.extract(buf.getvalue(), '.xlsx')[0]['location'])
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
        writer = PdfWriter()
        page = writer.add_blank_page(300, 300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 10 250 Td (Packing List: quantity 10 gross weight 120 kg) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        buf = io.BytesIO()
        writer.write(buf)
        self.assertIn('Packing List', document_text.extract(buf.getvalue(), '.pdf')[0]['text'])

    def test_blank_scanned_encrypted_and_overlimit_not_classified(self):
        from pypdf import PdfWriter
        for encrypted in [False, True]:
            writer = PdfWriter()
            writer.add_blank_page(300, 300)
            if encrypted:
                writer.encrypt('password')
            buf = io.BytesIO()
            writer.write(buf)
            with self.assertRaises(document_text.TextError):
                document_text.extract(buf.getvalue(), '.pdf')
        with self.assertRaises(document_text.TextError):
            document_text.extract(docx_bytes(), '.docx', max_chars=10)
        with self.assertRaises(document_text.TextError):
            document_text.extract(b'broken', '.docx')

    def test_pdf_with_empty_open_password_can_be_read(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
        for password in ['', 'required-password']:
            with self.subTest(password_required=bool(password)):
                writer = PdfWriter()
                page = writer.add_blank_page(300, 300)
                font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                    NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
                page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
                    DictionaryObject({NameObject('/F1'): font})})
                stream = DecodedStreamObject()
                stream.set_data(b'BT /F1 12 Tf 10 250 Td (Packing List: quantity 10 gross weight 120 kg) Tj ET')
                page[NameObject('/Contents')] = writer._add_object(stream)
                writer.encrypt(user_password=password, owner_password='owner-password')
                buf = io.BytesIO()
                writer.write(buf)
                if password:
                    with self.assertRaisesRegex(document_text.TextError, '需要打开密码'):
                        document_text.extract(buf.getvalue(), '.pdf')
                else:
                    sections = document_text.extract(buf.getvalue(), '.pdf')
                    self.assertIn('Packing List', sections[0]['text'])

    def test_pipeline_local_preview_and_write(self):
        from yida_archive.fetch_local import load as local_load, LocalClient
        root = Path(self.tmp.name) / 'source'
        target = root / '批次' / '附件'
        target.mkdir(parents=True)
        (target / '资料.docx').write_bytes(docx_bytes())
        instances = local_load(root)
        with patch.object(ai.Client, 'classify', return_value=(decision(), {'total_tokens': 30})) as call:
            preview = pipeline.run(self.cfg, instances, dry_run=True, verbose=False)
            self.assertEqual(preview.files[0].type_name, '箱单')
            self.assertFalse((Path(self.tmp.name) / '测试单位' / 'P1' / '编号记录.json').exists())
            written = pipeline.run(self.cfg, instances, client=LocalClient([root]), dry_run=False, verbose=False)
            self.assertEqual(call.call_count, 1)
        self.assertEqual(len(written.done), 1)
        with open(written.paths['csv'], encoding='utf-8-sig') as fh:
            row = next(csv.DictReader(fh))
        self.assertEqual(row['AI状态'], 'classified')
        self.assertEqual(row['AI缓存命中'], '是')
        self.assertEqual(Path(written.done[0][1]).read_bytes(), (target / '资料.docx').read_bytes())


    def test_web_preview_zip_save_cache_and_history_protection(self):
        from web import server as web
        body = docx_bytes()
        fresh = copy.deepcopy(web.STATE)
        fresh.update(batches=[{'id': 'b1', 'name': '合同审批', 'created': 1000,
            'files': [{'name': '资料.docx', 'size': len(body), 'uuid': 'u1',
                       'sha256': hashlib.sha256(body).hexdigest()}]}],
            blobs={'u1': body}, file_reviews={}, project_no='P1', project_secret='否')
        with patch.object(web, 'STATE', fresh), patch.object(web, 'SERVER_CONFIG', self.cfg.raw), \
             patch.object(web, 'OUTPUT_ROOT', self.tmp.name), \
             patch.object(ai.Client, 'classify', return_value=(decision(), {'total_tokens': 30})) as call:
            analyzed = web.analyze()
            self.assertEqual(analyzed['files'][0]['type'], '箱单')
            zipped, error = web.export_zip()
            self.assertFalse(error)
            with zipfile.ZipFile(io.BytesIO(zipped)) as z:
                filename = next(n for n in z.namelist() if n.endswith('.docx'))
                self.assertEqual(z.read(filename), body)
            self.assertTrue(web.save_result()['ok'])
            self.assertEqual(call.call_count, 1)
            for cache in (Path(self.tmp.name) / 'ai').glob('*.json'):
                cache.unlink()
            with patch.object(ai.Client, 'classify', side_effect=ai.AIError('模型不可用')):
                saved_again = web.save_result()
            self.assertTrue(saved_again['ok'])
            self.assertTrue(list((Path(self.tmp.name) / '测试单位' / 'P1').rglob('*.docx')))


if __name__ == '__main__':
    unittest.main()
