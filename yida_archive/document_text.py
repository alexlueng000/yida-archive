"""正文提取。只读取文字，不渲染图片，不执行宏或外部链接。"""
import io
import zipfile


SUPPORTED = {'.pdf', '.docx', '.xlsx', '.xls'}
IMAGES = {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tif', '.tiff', '.webp', '.heic', '.svg'}


class TextError(ValueError):
    pass


def extract(data, extension, max_chars=48000, max_units=120):
    """返回带定位信息的正文；超过界限整体拒绝，避免截断后误归档。"""
    ext = extension.lower()
    if ext not in SUPPORTED:
        raise TextError('首期不识别此格式的正文')
    if len(data) > 64 * 1024 * 1024:
        raise TextError('文档超过64MB识别上限')
    if ext in {'.docx', '.xlsx'}:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > 128 * 1024 * 1024:
                    raise TextError('文档解压内容过大')
        except zipfile.BadZipFile:
            raise TextError('文档格式损坏或已加密') from None
    sections = []
    total = 0

    def add(location, text):
        nonlocal total
        text = str(text or '').strip()
        if not text:
            return
        total += len(text)
        if total > max_chars:
            raise TextError('正文超过识别长度上限，未截断归档')
        sections.append({'location': location, 'text': text})

    try:
        if ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            # 部分PDF仅限制编辑/打印，打开密码为空，仍可正常读取正文。
            if reader.is_encrypted and not reader.decrypt(''):
                raise TextError('PDF需要打开密码，未能读取正文')
            if len(reader.pages) > max_units:
                raise TextError('PDF页数超过识别上限')
            for n, page in enumerate(reader.pages, 1):
                contents = page.get_contents()
                if contents and len(contents.get_data()) > 10 * 1024 * 1024:
                    raise TextError('PDF页面内容流过大')
                text = page.extract_text() or ''
                # 不让混合扫描文档凭少量文字被误认成已全文核验。
                if not text.strip() and page.get('/Resources', {}).get('/XObject'):
                    raise TextError('PDF含无正文的扫描或图片页面，首期不做OCR')
                add('第%d页' % n, text)
        elif ext == '.docx':
            from docx import Document
            doc = Document(io.BytesIO(data))
            for n, p in enumerate(doc.paragraphs, 1):
                add('段落%d' % n, p.text)
            for n, table in enumerate(doc.tables, 1):
                for r, row in enumerate(table.rows, 1):
                    add('表格%d行%d' % (n, r), ' | '.join(c.text for c in row.cells))
            for n, section in enumerate(doc.sections, 1):
                add('节%d页眉' % n, '\n'.join(p.text for p in section.header.paragraphs))
                add('节%d页脚' % n, '\n'.join(p.text for p in section.footer.paragraphs))
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > max_units:
                    raise TextError('工作表数量超过识别上限')
                for sheet in book.worksheets:
                    if sheet.max_row > 10000 or sheet.max_column > 200:
                        raise TextError('工作表范围超过识别上限')
                    for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if any(v is not None for v in row):
                            add('工作表%s第%d行' % (sheet.title, n),
                                ' | '.join('' if v is None else str(v) for v in row))
            finally:
                book.close()
        elif ext == '.xls':
            import xlrd
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                if book.nsheets > max_units:
                    raise TextError('工作表数量超过识别上限')
                for sheet in book.sheets():
                    if sheet.nrows > 10000 or sheet.ncols > 200:
                        raise TextError('工作表范围超过识别上限')
                    for n in range(sheet.nrows):
                        add('工作表%s第%d行' % (sheet.name, n + 1),
                            ' | '.join(str(v) for v in sheet.row_values(n)))
            finally:
                book.release_resources()
    except ImportError:
        raise TextError('缺少正文提取依赖，请安装requirements.txt') from None
    except TextError:
        raise
    except Exception:
        raise TextError('文档损坏、加密或正文无法解析') from None
    if sum(len(s['text']) for s in sections) < 20:
        raise TextError('未提取到足够正文；扫描件首期不做OCR')
    return sections
